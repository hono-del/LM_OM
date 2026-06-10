# -*- coding: utf-8 -*-
"""Generate bilingual Excel for over-implementation / mock-specific additions."""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import date

OUTPUT = r"system requirement\C-APP_過剰独自実装_開発者共有要件.xlsx"

ROWS = [
    {
        "no": 1,
        "type_ja": "過剰・独自",
        "type_en": "Over-implementation",
        "item_ja": "6シナリオ体験設計",
        "item_en": "6-Scenario Experience Design",
        "summary_ja": "納車・充電・OTA・日常・トラブル・販売店の6シナリオを横断し、通知・チャット・警告灯を一連の体験として接続するデモ演出。",
        "summary_en": "Cross-cutting demo that connects notifications, chat, and warning lamps across six owner scenarios.",
        "mock_ja": "demo-data.js の固定JSONと demo-runner.js によるクライアント側シナリオ切替。通知既読・警告灯状態は sessionStorage。",
        "mock_en": "Client-side scenario switching via fixed JSON in demo-data.js and demo-runner.js. Read/ warning state in sessionStorage.",
        "prod_ja": "バックエンドの通知配信・ルールエンジン(7.7)・車両イベントに基づく実データ駆動。ハードコードされたシナリオランナーは本番に持ち込まない。",
        "prod_en": "Drive from backend notifications, rules engine (7.7), and vehicle events. Do not carry hardcoded scenario runner into production.",
        "files": "owners-assistant/demo-data.js, demo-runner.js",
        "wbs_ja": "MVP WBS に同名項目なし。通知(7.x)・チャット(4.x)・警告灯(6.x)の横断デモとして追加。",
        "wbs_en": "No matching WBS item in MVP. Added as cross-cutting demo spanning Notifications, Chat, and Warning Lamps.",
        "decision_ja": "本番に取り込むか／デモ専用に留めるか／シナリオ数を削減するかを決定してください。",
        "decision_en": "Decide adopt in production, demo-only, or reduce scenario count.",
        "rec_ja": "プレゼン・UAT用デモとして維持。本番はAPI連携後にシナリオエンジン要否を再評価。",
        "rec_en": "Keep for presentation/UAT. Re-evaluate scenario engine after API integration.",
        "priority": "Medium",
        "notes": "WBS外の独自追加",
    },
    {
        "no": 2,
        "type_ja": "過剰・独自",
        "type_en": "Over-implementation",
        "item_ja": "初回オンボーディング 7ステップ（画像付き）",
        "item_en": "7-Step First-Time Onboarding (with images)",
        "summary_ja": "初回ログイン後の車両オーナー向けガイド。各ステップに画像・説明文。スキップと全完了を区別し、全完了時に「ガイド完了」バッジを付与。",
        "summary_en": "First-time owner guide after login with images and copy per step. Distinguishes skip vs full completion; badge on full completion only.",
        "mock_ja": "C-APPシェル外の独立HTML（onboarding.html）。7画面・文言は onboarding-data.js の静的定義。通過/完了フラグは localStorage（capp_has_completed_onboarding / capp_onboarding_fully_completed）。",
        "mock_en": "Standalone HTML outside C-APP shell. 7 screens; copy in static onboarding-data.js. Pass/full-complete flags in localStorage.",
        "prod_ja": "C-APPアプリ内の画面フロー（React Native / WebView 等）に統合。ステップ定義・画像はCMS/API管理。完了状態はユーザープロファイルAPIで永続化（端末ローカル保存に依存しない）。スキップ/全完了をサーバー側でも区別。",
        "prod_en": "Integrate as in-app screen flow. Step content/images via CMS/API. Persist completion on user profile API (not device localStorage). Distinguish skip vs full complete server-side.",
        "files": "owners-assistant/onboarding.html, onboarding-data.js, onboarding-status.js",
        "wbs_ja": "5.3 / 7.9 / 9.x に近いが、独立フローとして厚め。Excel上は分散した要件。",
        "wbs_en": "Related to 5.3 / 7.9 / 9.x but thicker standalone flow; requirements scattered in Excel.",
        "decision_ja": "ステップ数・画像要件・C-APPシェル内統合・状態管理API仕様を決定してください。",
        "decision_en": "Decide step count, images, in-app integration, and state-management API spec.",
        "rec_ja": "MVPは3〜5ステップに縮小。本番はユーザープロファイルAPI＋CMSで管理。",
        "rec_en": "Reduce to 3–5 steps for MVP. Manage via user profile API + CMS in production.",
        "priority": "High",
        "notes": "localStorageはモック限定",
    },
    {
        "no": 3,
        "type_ja": "過剰・独自",
        "type_en": "Over-implementation",
        "item_ja": "警告灯：2件限定＋多段オーバーレイ",
        "item_en": "Warning Lamps: 2 Active Cases + Multi-Step Overlays",
        "summary_ja": "点灯中の警告灯を強調表示し、詳細から共有・入庫予約・ロードアシスタンス等へ誘導。全警告灯一覧は別画面。",
        "summary_en": "Highlight active warning lamps; detail flows to share, service booking, roadside assistance, etc. Full lamp list on separate screen.",
        "mock_ja": "WARNING_CASES の静的データ。警告灯状態は sessionStorage。オーバーレイはクライアント側モーダルで多段遷移（外部連携なし）。",
        "mock_en": "Static WARNING_CASES data. Lamp state in sessionStorage. Multi-step overlays as client-side modals without real integrations.",
        "prod_ja": "車両テレマティクス/APIからのリアルタイム警告灯データ(6.1/6.2)。予約・ロードアシスタンスは外部サービス連携またはディープリンク。表示件数・フェーズはMVP/P1で切分。",
        "prod_en": "Real-time lamp data from vehicle telematics/API (6.1/6.2). Booking/roadside via external services or deep links. Phase-split MVP/P1.",
        "files": "owners-assistant/demo-runner.js, demo-data.js, warning-lamps.js",
        "wbs_ja": "6.1/6.2よりPhase 2–3相当の体験を先行実装。ロードアシスタンス・販売店予約はWBS上後期。",
        "wbs_en": "Phase 2–3 experience ahead of 6.1/6.2. Roadside assistance and booking are later WBS items.",
        "decision_ja": "オーバーレイ各段階の本番スコープ（MVP/P1/P2）と連携先APIを切り分けてください。",
        "decision_en": "Split production scope (MVP/P1/P2) and integration APIs per overlay step.",
        "rec_ja": "MVPは6.1グリッド＋6.2詳細まで。予約・ロードアシスタンスはP1以降。",
        "rec_en": "MVP: 6.1 grid + 6.2 detail. Booking/roadside assistance from P1 onward.",
        "priority": "High",
        "notes": "ブレーキ警告は入庫先オートバックス（モック）",
    },
    {
        "no": 4,
        "type_ja": "過剰・独自",
        "type_en": "Over-implementation",
        "item_ja": "ホーム 2×2 お知らせカード＋動的ヒーロー履歴",
        "item_en": "Home 2×2 Notification Cards + Dynamic Hero History",
        "summary_ja": "ホームに4枚のお知らせカード（2×2）と、ユーザー行動に応じたヒーロー領域の表示切替。",
        "summary_en": "Four notification cards (2×2) on Home and a hero area that changes based on user engagement.",
        "mock_ja": "HOME_NOTIFICATION_CARDS の固定データ。ヒーロー完了履歴は localStorage（capp_hero_completed）。",
        "mock_en": "Fixed HOME_NOTIFICATION_CARDS data. Hero completion history in localStorage (capp_hero_completed).",
        "prod_ja": "通知API/CMS(7.5)とルールエンジン(7.7)によるパーソナライズ配信。閲覧・完了状態はサーバー側で管理。プッシュ通知(7.4)と連携。",
        "prod_en": "Personalized delivery via notification API/CMS (7.5) and rules engine (7.7). View/completion state on server. Integrate push (7.4).",
        "files": "owners-assistant/demo-data.js, app.html, demo-runner.js",
        "wbs_ja": "7.5 Smart Alert Cards の具体化。ルールエンジン(7.7)・プッシュ(7.4)は未実装。",
        "wbs_en": "Concrete UI for 7.5 Smart Alert Cards. Rules engine (7.7) and push (7.4) not in mock.",
        "decision_ja": "カード枚数・レイアウト・CMS/ルールエンジン連携仕様を確定してください。",
        "decision_en": "Confirm card count, layout, and CMS/rules-engine integration spec.",
        "rec_ja": "UI参考として採用可。本番はAPI/CMS駆動に全面置換。",
        "rec_en": "Adopt as UI reference. Fully replace with API/CMS in production.",
        "priority": "Medium",
        "notes": "",
    },
    {
        "no": 5,
        "type_ja": "過剰・独自",
        "type_en": "Over-implementation",
        "item_ja": "設定：デモ用リセット機能",
        "item_en": "Settings: Demo Reset Utilities",
        "summary_ja": "デモ再実行のため、初回導線・お知らせ・チャット履歴・ステータス進捗を初期化する管理機能。",
        "summary_en": "Admin utilities to reset first-login flow, notifications, chat history, and status progress for demo reruns.",
        "mock_ja": "設定画面から localStorage / sessionStorage のキーを直接削除（onboarding-status.js, chat-history.js, user-progress.js 等）。",
        "mock_en": "Settings UI directly clears localStorage/sessionStorage keys (onboarding-status.js, chat-history.js, user-progress.js, etc.).",
        "prod_ja": "本番アプリには実装しない。必要なら開発・ステージング環境のみ feature flag で提供。本番のユーザーデータ削除はアカウント管理・サポート手順に従う。",
        "prod_en": "Do not ship in production app. If needed, dev/staging only via feature flag. Production data deletion follows account/support procedures.",
        "files": "owners-assistant/onboarding-status.js, app-settings.js, chat-history.js, user-progress.js",
        "wbs_ja": "10.1 Settingsの一部。本番向けリセットは通常スコープ外（QA/デモ用）。",
        "wbs_en": "Subset of 10.1 Settings. Production reset utilities typically out of scope.",
        "decision_ja": "本番ビルドに含めるか、開発/ステージングのみに限定するか。",
        "decision_en": "Include in production builds or restrict to dev/staging only.",
        "rec_ja": "本番対象外。開発・デモ環境用フラグで制御推奨。",
        "rec_en": "Exclude from production. Control via dev/demo feature flags.",
        "priority": "Low",
        "notes": "デモ運用必須・本番非搭載",
    },
    {
        "no": 6,
        "type_ja": "過剰・独自",
        "type_en": "Over-implementation",
        "item_ja": "別ドキュメント manual.html（本格ビューア＋PDF）",
        "item_en": "Standalone manual.html (Full Viewer + PDF)",
        "summary_ja": "車両マニュアルの章立て閲覧とPDFダウンロード。チャット・FAQからの参照先。",
        "summary_en": "Chapter-based vehicle manual viewer and PDF download. Linked from chat and FAQ.",
        "mock_ja": "C-APPシェル外の独立HTML。章データは data/manual-data.js の静的JSON。チャットからは別タブ/外部リンクで遷移。",
        "mock_en": "Standalone HTML outside C-APP shell. Chapters in static manual-data.js. Chat opens via separate tab/external link.",
        "prod_ja": "C-APP内タブまたはWebViewに統合。マニュアルコンテンツはCMS/API(3.2)で配信。チャット・FAQからのディープリンクをアプリ内ナビで一貫化。",
        "prod_en": "Integrate as in-app tab or WebView. Manual content via CMS/API (3.2). Consistent in-app deep links from chat/FAQ.",
        "files": "manual.html, manual-app.js, data/manual-data.js",
        "wbs_ja": "3.2 Digital Manual のプロトタイプがアプリ外に存在。アプリ内統合は未完了。",
        "wbs_en": "3.2 Digital Manual prototype outside app. In-app integration not done.",
        "decision_ja": "C-APP内統合方式、CMS連携、ディープリンク仕様を決定してください。",
        "decision_en": "Decide in-app integration approach, CMS linkage, and deep-link spec.",
        "rec_ja": "最優先でC-APPシェル内に統合（ギャップ分析 推奨#1）。",
        "rec_en": "Prioritize integration into C-APP shell (Gap Analysis Recommendation #1).",
        "priority": "High",
        "notes": "静的JSONはモック限定",
    },
    {
        "no": 7,
        "type_ja": "スコープ外",
        "type_en": "Out of Scope",
        "item_ja": "lean3-pages-editor.html 等（編集ツール）",
        "item_en": "lean3-pages-editor.html etc. (Editor Tools)",
        "summary_ja": "Lean3ページ比較・編集用の社内ツール。C-APP製品機能ではない。",
        "summary_en": "Internal tools for Lean3 page comparison/editing. Not a C-APP product feature.",
        "mock_ja": "リポジトリ直下の静的HTMLツール。ローカルで直接開いて使用。",
        "mock_en": "Static HTML tools in repo root. Opened locally for internal use.",
        "prod_ja": "C-APP本番ビルド・配布物・モバイルアプリに含めない。必要なら別リポジトリで管理。",
        "prod_en": "Exclude from C-APP production build, deliverables, and mobile app. Manage in separate repo if needed.",
        "files": "lean3-pages-editor.html, lean3-comparison*.html（リポジトリ直下）",
        "wbs_ja": "C-APP Effort EstimationのWBSに該当なし。",
        "wbs_en": "Not in C-APP Effort Estimation WBS.",
        "decision_ja": "本番リポジトリから除外するか、別リポジトリへ移管するか。",
        "decision_en": "Remove from production repo or move to separate repository.",
        "rec_ja": "C-APP本番ビルド・配布物から除外。",
        "rec_en": "Exclude from C-APP production build and deliverables.",
        "priority": "Low",
        "notes": "開発者向けツール",
    },
    {
        "no": 8,
        "type_ja": "Phase 1 先行",
        "type_en": "Phase 1 Early",
        "item_ja": "4.3 Home Search → Chat",
        "item_en": "4.3 Home Search → Chat",
        "summary_ja": "ホームの検索からチャットタブへ遷移し、検索語をチャット入力に引き継ぐ。",
        "summary_en": "Home search navigates to Chat tab and passes the query into chat input.",
        "mock_ja": "app.html 内のクライアント側タブ切替と入力欄への文字列セット。バックエンド呼び出しなし。",
        "mock_en": "Client-side tab switch and input prefill in app.html. No backend call.",
        "prod_ja": "アプリの標準ナビゲーション（React Navigation 等）で実装。検索イベントの分析ログ送信を検討。AIチャットAPI(4.6)への接続は別途。",
        "prod_en": "Implement via standard app navigation. Consider analytics for search events. AI chat API (4.6) connection is separate.",
        "files": "owners-assistant/app.html",
        "wbs_ja": "Excel上はPhase 1。モックに先行実装済み。",
        "wbs_en": "Phase 1 in Excel; already implemented in mock.",
        "decision_ja": "MVPに含めるかPhase 1のまま維持するか。実装を参考に工数見積を更新するか。",
        "decision_en": "Include in MVP or keep as Phase 1. Update effort estimate using mock as reference.",
        "rec_ja": "UX参考として維持。本番ナビゲーション・分析要件を別途定義。",
        "rec_en": "Keep as UX reference. Define production navigation and analytics separately.",
        "priority": "Medium",
        "notes": "",
    },
    {
        "no": 9,
        "type_ja": "Phase 1 先行",
        "type_en": "Phase 1 Early",
        "item_ja": "5.2 / 7.3 等 Phase 1 機能の先行実装",
        "item_en": "5.2 / 7.3 etc. Phase 1 Features Implemented Early",
        "summary_ja": "チャット内警告灯グリッド(5.2)、関連通知UI(7.3)など、Phase 1相当の画面・フローを先行搭載。",
        "summary_en": "Phase 1 screens/flows such as in-chat warning lamp grid (5.2) and related notification UI (7.3).",
        "mock_ja": "静的デモデータとクライアント側UIのみ。Helpful/Not Helpful・チャット履歴は localStorage。API・AI本番接続なし。",
        "mock_en": "Static demo data and client UI only. Helpful/Not Helpful and chat history in localStorage. No API/AI production connection.",
        "prod_ja": "各WBS項目の正式API・バックエンドと接続。チャット履歴・評価ログはサーバー永続化。MVPに含めるかPhase 1に戻すかをスコープで確定。",
        "prod_en": "Connect to formal APIs/backends per WBS item. Persist chat history and feedback on server. Confirm MVP vs Phase 1 scope.",
        "files": "owners-assistant/warning-lamps.js, chat-history.js, app.html",
        "wbs_ja": "Phase 1だがMVPモックに含まれる。工数表とのフェーズ差異に注意。",
        "wbs_en": "Phase 1 but present in MVP mock. Mind phase mismatch vs effort sheet.",
        "decision_ja": "先行実装をMVP正式要件に昇格するか、Phase 1実装計画に戻すか。API接続仕様を確定してください。",
        "decision_en": "Promote to formal MVP or revert to Phase 1 plan. Confirm API integration spec.",
        "rec_ja": "UI参考として活用。localStorage依存部分は本番APIに置換必須。",
        "rec_en": "Use as UI reference. Must replace localStorage-dependent parts with production APIs.",
        "priority": "Medium",
        "notes": "localStorageはモック限定",
    },
]

COLUMNS = [
    ("No", 5),
    ("区分 / Type", 14),
    ("項目（日本語）", 26),
    ("Item (English)", 26),
    ("機能概要（日本語）", 34),
    ("Functional Summary (EN)", 34),
    ("【モック】実装（現状・参考）", 36),
    ("【Mock】Implementation (Current)", 36),
    ("【本番】実装方針", 36),
    ("【Production】Target", 36),
    ("該当ファイル", 28),
    ("Excel WBS との関係（日本語）", 32),
    ("WBS Relationship (English)", 32),
    ("開発者への判断依頼（日本語）", 32),
    ("Developer Decision Required (EN)", 32),
    ("推奨方針（日本語）", 28),
    ("Recommended Action (EN)", 28),
    ("優先度 / Priority", 10),
    ("備考 / Notes", 20),
]

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
MOCK_HEADER_FILL = PatternFill("solid", fgColor="5B7EA8")
PROD_HEADER_FILL = PatternFill("solid", fgColor="2E6B4A")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=9)
TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
META_FONT = Font(size=10, color="333333")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TYPE_FILLS = {
    "過剰・独自": PatternFill("solid", fgColor="FFF2CC"),
    "スコープ外": PatternFill("solid", fgColor="F4CCCC"),
    "Phase 1 先行": PatternFill("solid", fgColor="D9EAD3"),
}

MOCK_FILL = PatternFill("solid", fgColor="E8F0F8")
PROD_FILL = PatternFill("solid", fgColor="E8F5E9")


def style_header_row(ws, row, ncol, fills=None):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        if fills and c in fills:
            cell.fill = fills[c]
        else:
            cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def autosize_and_freeze(ws, header_row, ncol):
    for i, (_, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncol)}{ws.max_row}"


def build_main_sheet(wb):
    ws = wb.create_sheet("Over-Implementation List")
    ncol = len(COLUMNS)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    t = ws.cell(1, 1, "C-APP モック：過剰・独自実装 — 開発者共有要件 / Mock Over-Implementation & Additions")
    t.font = TITLE_FONT
    t.alignment = Alignment(vertical="center")

    meta = [
        f"作成日 / Created: {date.today().isoformat()}",
        "参照 / Reference: C-APP_モック_vs_EffortEstimation_ギャップ分析.md §4",
        "比較元 / Source WBS: CMC NDG Next - C-App Effort Estimation (R-Edge).xlsx",
        "対象モック / Mock: owners-assistant/, manual.html（静的HTML + Vanilla JS）",
        "※【モック】列は現行プロトタイプの実装手段。【本番】列は本番開発で採用すべき方針。混同しないこと。",
        "※【Mock】= prototype implementation. 【Production】= target for production. Do not treat mock details as production spec.",
    ]
    for i, line in enumerate(meta, 2):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=ncol)
        c = ws.cell(i, 1, line)
        c.font = META_FONT

    header_row = 8
    mock_cols = {7, 8}
    prod_cols = {9, 10}
    header_fills = {}
    for c in range(1, ncol + 1):
        if c in mock_cols:
            header_fills[c] = MOCK_HEADER_FILL
        elif c in prod_cols:
            header_fills[c] = PROD_HEADER_FILL

    for i, (name, _) in enumerate(COLUMNS, 1):
        ws.cell(header_row, i, name)

    style_header_row(ws, header_row, ncol, header_fills)

    for r_idx, row in enumerate(ROWS, header_row + 1):
        values = [
            row["no"],
            f"{row['type_ja']} / {row['type_en']}",
            row["item_ja"],
            row["item_en"],
            row["summary_ja"],
            row["summary_en"],
            row["mock_ja"],
            row["mock_en"],
            row["prod_ja"],
            row["prod_en"],
            row["files"],
            row["wbs_ja"],
            row["wbs_en"],
            row["decision_ja"],
            row["decision_en"],
            row["rec_ja"],
            row["rec_en"],
            row["priority"],
            row["notes"],
        ]
        type_fill = TYPE_FILLS.get(row["type_ja"])
        for c_idx, val in enumerate(values, 1):
            cell = ws.cell(r_idx, c_idx, val)
            cell.alignment = WRAP
            cell.border = BORDER
            if type_fill and c_idx <= 2:
                cell.fill = type_fill
            elif c_idx in mock_cols:
                cell.fill = MOCK_FILL
            elif c_idx in prod_cols:
                cell.fill = PROD_FILL
        ws.row_dimensions[r_idx].height = 90

    autosize_and_freeze(ws, header_row, ncol)
    ws.row_dimensions[header_row].height = 36


def build_summary_sheet(wb):
    ws = wb.create_sheet("Overview", 0)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 58
    ws.column_dimensions["C"].width = 58

    lines = [
        ("文書目的 / Purpose",
         "モックに意図的に追加された過剰・独自実装を開発者へ共有し、本番スコープへの採用・縮小・除外を判断する。",
         "Share mock-specific additions so developers can decide adopt, reduce, or exclude from production."),
        ("モック vs 本番の読み方",
         "【モック】列＝現行HTML/JSプロトタイプの実装（localStorage・静的JSON・独立HTML等）。【本番】列＝本番で採用すべき方針。モックの技術選択をそのまま本番要件にしない。",
         "【Mock】= current prototype (localStorage, static JSON, standalone HTML). 【Production】= production target. Do not copy mock tech choices as production requirements."),
        ("対象読者 / Audience",
         "C-APP フロントエンド・バックエンド開発者、PM、QA",
         "C-APP frontend/backend developers, PM, QA"),
        ("凡例：過剰・独自",
         "Excel MVPスコープにない、またはデモ厚みが本番要件を超える実装",
         "Not in Excel MVP scope, or demo richness exceeds production requirements"),
        ("凡例：スコープ外",
         "C-APP製品とは無関係のツール・ファイル",
         "Tools/files unrelated to C-APP product"),
        ("凡例：Phase 1 先行",
         "工数表上Phase 1だがモックに先行実装済み",
         "Phase 1 in effort sheet but already in mock"),
        ("記入・回答方法",
         "「開発者への判断依頼」列に対し、採用/縮小/本番対象外/Phase維持 をチームで記録。",
         "Record adopt / reduce / exclude / keep phase in Developer Decision column."),
        ("関連ドキュメント",
         "system requirement/C-APP_モック_vs_EffortEstimation_ギャップ分析.md",
         "Gap analysis MD (section 4 & 7)"),
        ("次アクション（推奨）",
         "1) manual.htmlをC-APP内統合 2) オンボーディングはAPI/CMS化 3) localStorage依存の排除 4) デモリセットは本番除外",
         "1) Integrate manual.html 2) Onboarding via API/CMS 3) Remove localStorage deps 4) Exclude demo resets from prod"),
    ]

    ws.cell(1, 1, "C-APP 過剰・独自実装 — 概要").font = TITLE_FONT
    ws.merge_cells("A1:C1")

    headers = ["項目 / Item", "日本語", "English"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(3, i, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER

    for r, (a, b, c) in enumerate(lines, 4):
        for col, val in enumerate([a, b, c], 1):
            cell = ws.cell(r, col, val)
            cell.alignment = WRAP
            cell.border = BORDER
        ws.row_dimensions[r].height = 52


def build_response_sheet(wb):
    ws = wb.create_sheet("Developer Response")
    cols = [
        ("No", 5),
        ("項目 / Item", 30),
        ("判断 / Decision", 18),
        ("判断（日本語メモ）", 36),
        ("Decision Notes (English)", 36),
        ("担当 / Owner", 16),
        ("期限 / Due", 12),
        ("ステータス / Status", 14),
    ]
    for i, (name, w) in enumerate(cols, 1):
        ws.cell(1, i, name)
        ws.column_dimensions[get_column_letter(i)].width = w
    style_header_row(ws, 1, len(cols))
    ws.freeze_panes = "A2"

    decisions = "採用 Adopt | 縮小 Reduce | 本番対象外 Exclude | Phase維持 Keep Phase | 要検討 TBD"
    for r, row in enumerate(ROWS, 2):
        ws.cell(r, 1, row["no"])
        ws.cell(r, 2, f"{row['item_ja']} / {row['item_en']}")
        ws.cell(r, 3, decisions)
        for c in range(1, len(cols) + 1):
            ws.cell(r, c).alignment = WRAP
            ws.cell(r, c).border = BORDER


def main():
    wb = Workbook()
    wb.remove(wb.active)
    build_summary_sheet(wb)
    build_main_sheet(wb)
    build_response_sheet(wb)
    wb.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    main()
